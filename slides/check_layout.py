"""スライドのレイアウト検査。

LibreOffice が無い環境向け。座標とテキスト量から以下を機械的に検査する。

  1. スライド境界のはみ出し
  2. 左右マージン不足
  3. テキストの推定必要高さ > 枠の高さ（＝はみ出しの疑い）
  4. テキストボックス同士の重なり

日本語は全角なので、フォントサイズ ≒ 1文字の幅として概算する。
"""
import math
import sys
import unicodedata

from pptx import Presentation
from pptx.util import Emu

EMU_PER_IN = 914400
SW, SH = 10.0, 5.625
MARGIN_MIN = 0.40


def inch(v) -> float:
    return v / EMU_PER_IN if v is not None else 0.0


def text_width_pt(text: str, size_pt: float) -> float:
    """文字列の表示幅を pt で概算する。全角=1.0em、半角=0.55em。"""
    w = 0.0
    for ch in text:
        ea = unicodedata.east_asian_width(ch)
        w += size_pt * (1.0 if ea in ("W", "F", "A") else 0.55)
    return w


def estimate_text_height(tf, box_w_in: float) -> float:
    """テキストフレームの必要高さを inch で概算する。"""
    box_w_pt = box_w_in * 72
    if box_w_pt <= 0:
        return 0.0
    total_pt = 0.0
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs)
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        size = max(sizes) if sizes else 12.0
        before = p.space_before.pt if p.space_before is not None else 0.0
        after = p.space_after.pt if p.space_after is not None else 0.0
        if not text:
            total_pt += size * 0.6 + before + after
            continue
        # 明示的な改行も行数に数える
        lines = 0
        for seg in text.split("\n"):
            w = text_width_pt(seg, size)
            lines += max(1, math.ceil(w / box_w_pt - 1e-9))
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        total_pt += lines * size * ls + before + after
    return total_pt / 72


def overlap(a, b) -> float:
    """2矩形の重なり面積（平方インチ）。"""
    ax1, ay1, ax2, ay2 = a
    bx1, by1, bx2, by2 = b
    dx = min(ax2, bx2) - max(ax1, bx1)
    dy = min(ay2, by2) - max(ay1, by1)
    return dx * dy if dx > 0 and dy > 0 else 0.0


def main(path: str) -> int:
    prs = Presentation(path)
    problems = 0

    for idx, slide in enumerate(prs.slides, start=1):
        issues = []
        text_boxes = []  # (name, rect, tf)

        for sh in slide.shapes:
            x, y = inch(sh.left), inch(sh.top)
            w, h = inch(sh.width), inch(sh.height)
            r = (x, y, x + w, y + h)
            kind = sh.shape_type
            label = f"{sh.shape_id}:{kind}"

            # 1. スライド外
            if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
                issues.append(
                    f"スライド外 {label} "
                    f"x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}"
                )

            is_bg = w >= SW - 0.02 and h >= SH - 0.02
            if not is_bg:
                # 2. マージン
                if 0 <= x < MARGIN_MIN - 0.01:
                    issues.append(f"左マージン不足 {label} x={x:.2f}")
                if x + w > SW - MARGIN_MIN + 0.01:
                    issues.append(f"右マージン不足 {label} right={x + w:.2f}")

            if sh.has_text_frame and sh.text_frame.text.strip():
                need = estimate_text_height(sh.text_frame, w)
                # 3. 高さ不足（1.15 倍の余裕を見る）
                if need > h * 1.15 + 0.02:
                    preview = sh.text_frame.text.strip().replace("\n", " ")[:34]
                    issues.append(
                        f"高さ不足の疑い {label} 必要≈{need:.2f}in / 枠{h:.2f}in "
                        f"「{preview}」"
                    )
                text_boxes.append((label, r, sh))

        # 4. テキストボックス同士の重なり
        for i in range(len(text_boxes)):
            for j in range(i + 1, len(text_boxes)):
                n1, r1, _ = text_boxes[i]
                n2, r2, _ = text_boxes[j]
                ov = overlap(r1, r2)
                if ov > 0.02:
                    issues.append(f"テキスト重なり {n1} × {n2} 面積={ov:.3f}in²")

        if issues:
            problems += len(issues)
            print(f"\n── スライド {idx} ──")
            for it in issues:
                print(f"  ! {it}")

    print()
    if problems == 0:
        print("✓ 検出された問題なし")
    else:
        print(f"⚠ 合計 {problems} 件")
    return problems


if __name__ == "__main__":
    sys.exit(0 if main(sys.argv[1] if len(sys.argv) > 1
                       else "phase4_frontend.pptx") == 0 else 1)
